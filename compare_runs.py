"""Backtest sweep across 27 configurations for the BTC 5-minute up/down bot.

Runs the strategy engine over historical 1-minute candles for every combination
of 9 confidence thresholds x 3 sizing modes (27 configs) and writes an Excel
workbook (openpyxl) with three sheets:

* ``Summary``            - one row per config, sorted by final bankroll.
* ``Best Config Trades`` - the trade log of the best-performing config.
* ``Bankroll Curves``    - per-config bankroll after each window (configs as
                           columns).

Candle format (shared): ``dict`` with keys ``ts`` (int, ms epoch of candle
open), ``open``, ``high``, ``low``, ``close``, ``volume`` (floats).

Simulation model per 5-minute window
-------------------------------------
A window is 5 consecutive 1-minute candles aligned to a 300s boundary. The
decision is taken at T-10s before the window closes; at that point the first
four candles (offsets 0..180s) have closed and the fifth (offset 240s) is still
in progress. We therefore feed the strategy the closed candles up to T-10s and
use the last closed candle's close as the current price. Resolution compares the
window's final candle close against ``window_open`` (the price to beat):
``up_wins = close >= window_open``.
"""

from __future__ import annotations

import argparse
import json
import logging
import os

from openpyxl import Workbook

import strategy

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

THRESHOLDS = [0.0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8]
SIZING_MODES = ["flat", "safe", "aggressive"]

_MINUTE_MS = 60_000
_WINDOW_MS = 300_000
_DECISION_OFFSET_MS = 290_000   # T-10s within the 300s window
_LAST_CLOSED_OFFSET_MS = 180_000  # close of the 4th candle (index 3)


def _env_float(name: str, default: float) -> float:
    try:
        return float(os.environ.get(name, default))
    except (TypeError, ValueError):
        return default


def _stake_for(mode: str, bankroll: float, starting: float, min_bet: float) -> float:
    """Compute the stake for a trade given the sizing mode.

    ``flat``       - fixed MIN_BET per trade.
    ``safe``       - 25% of bankroll (at least MIN_BET).
    ``aggressive`` - bet only profit above the starting bankroll; if none, MIN_BET.
    """
    if mode == "flat":
        stake = min_bet
    elif mode == "safe":
        stake = max(min_bet, 0.25 * bankroll)
    elif mode == "aggressive":
        stake = max(bankroll - starting, min_bet)
    else:
        raise ValueError(f"Unknown sizing mode: {mode}")
    # Never stake more than the available bankroll.
    return min(stake, bankroll)


def build_windows(candles: list[dict]) -> list[dict]:
    """Group candles into simulatable 5-minute windows.

    Returns a list of window descriptors sorted by ``window_ts``, each with the
    window's five candles present, the pre-decision candle history, the decision
    price, and the resolution close.
    """
    by_ts = {c["ts"]: c for c in candles}
    sorted_candles = [by_ts[ts] for ts in sorted(by_ts)]
    if not sorted_candles:
        return []

    windows = []
    seen_ts = set(by_ts)
    for c in sorted_candles:
        ts = c["ts"]
        if ts % _WINDOW_MS != 0:
            continue  # only candles aligned to a window boundary start a window
        offsets = [ts + i * _MINUTE_MS for i in range(5)]
        if not all(o in seen_ts for o in offsets):
            continue
        window_candles = [by_ts[o] for o in offsets]
        window_open = window_candles[0]["open"]
        # Candles closed by T-10s: history up to and including offset +180s.
        decision_cutoff = ts + _LAST_CLOSED_OFFSET_MS
        history = [x for x in sorted_candles if x["ts"] <= decision_cutoff]
        current_price = by_ts[decision_cutoff]["close"]
        resolution_close = window_candles[4]["close"]
        windows.append({
            "window_ts": ts,
            "window_open": window_open,
            "current_price": current_price,
            "history": history,
            "resolution_close": resolution_close,
            "up_wins": resolution_close >= window_open,
        })
    return windows


def simulate_config(windows: list[dict], threshold: float, mode: str,
                    starting: float, min_bet: float) -> dict:
    """Run one (threshold, sizing mode) configuration over all windows."""
    bankroll = starting
    peak = starting
    max_drawdown = 0.0
    trades = 0
    wins = 0
    total_pnl = 0.0
    busts = 0
    trade_log: list[dict] = []
    curve: list[float] = []

    for w in windows:
        signal = strategy.analyze(w["history"], w["window_open"],
                                  w["current_price"], ticks=None)
        if signal["confidence"] >= threshold:
            delta_pct = (w["current_price"] - w["window_open"]) / w["window_open"] * 100.0
            entry_price = strategy.estimate_token_price(delta_pct)
            stake = _stake_for(mode, bankroll, starting, min_bet)
            shares = stake / entry_price
            side = signal["side"]
            won = (side == "UP" and w["up_wins"]) or (side == "DOWN" and not w["up_wins"])
            if won:
                pnl = shares * (1.0 - entry_price)
                wins += 1
            else:
                pnl = -shares * entry_price
            bankroll += pnl
            total_pnl += pnl
            trades += 1
            trade_log.append({
                "window_ts": w["window_ts"],
                "side": side,
                "confidence": signal["confidence"],
                "score": signal["score"],
                "delta_pct": delta_pct,
                "entry_price": entry_price,
                "stake": stake,
                "shares": shares,
                "outcome": "WIN" if won else "LOSS",
                "pnl": pnl,
                "bankroll_after": bankroll,
            })
            # Bankroll floor: reset to starting on a bust.
            if bankroll < min_bet:
                busts += 1
                bankroll = starting

        # Drawdown tracking on the running equity curve.
        if bankroll > peak:
            peak = bankroll
        if peak > 0:
            dd = (peak - bankroll) / peak
            if dd > max_drawdown:
                max_drawdown = dd
        curve.append(bankroll)

    win_rate = (wins / trades) if trades else 0.0
    return {
        "threshold": threshold,
        "mode": mode,
        "label": f"conf{threshold:.1f}_{mode}",
        "trades": trades,
        "wins": wins,
        "win_rate": win_rate,
        "total_pnl": total_pnl,
        "final_bankroll": bankroll,
        "max_drawdown": max_drawdown,
        "busts": busts,
        "trade_log": trade_log,
        "curve": curve,
    }


def run_all(candles: list[dict], starting: float, min_bet: float) -> tuple[list[dict], list[dict]]:
    """Run all 27 configurations. Returns (results, windows)."""
    windows = build_windows(candles)
    logger.info("Built %d simulatable 5-minute windows", len(windows))
    results = []
    for threshold in THRESHOLDS:
        for mode in SIZING_MODES:
            results.append(simulate_config(windows, threshold, mode, starting, min_bet))
    return results, windows


def write_workbook(results: list[dict], windows: list[dict], output: str) -> None:
    """Write the 3-sheet Excel workbook."""
    ranked = sorted(results, key=lambda r: r["final_bankroll"], reverse=True)

    wb = Workbook()

    # --- Sheet 1: Summary -------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    summary_headers = [
        "threshold", "sizing_mode", "trades", "wins", "win_rate",
        "total_pnl", "final_bankroll", "max_drawdown", "busts",
    ]
    ws.append(summary_headers)
    for r in ranked:
        ws.append([
            r["threshold"], r["mode"], r["trades"], r["wins"],
            round(r["win_rate"], 4), round(r["total_pnl"], 6),
            round(r["final_bankroll"], 6), round(r["max_drawdown"], 6),
            r["busts"],
        ])

    # --- Sheet 2: Best Config Trades -------------------------------------
    ws2 = wb.create_sheet("Best Config Trades")
    best = ranked[0] if ranked else None
    trade_headers = [
        "window_ts", "side", "confidence", "score", "delta_pct",
        "entry_price", "stake", "shares", "outcome", "pnl", "bankroll_after",
    ]
    ws2.append(trade_headers)
    if best is not None:
        for t in best["trade_log"]:
            ws2.append([
                t["window_ts"], t["side"], round(t["confidence"], 4),
                round(t["score"], 4), round(t["delta_pct"], 6),
                round(t["entry_price"], 6), round(t["stake"], 6),
                round(t["shares"], 6), t["outcome"], round(t["pnl"], 6),
                round(t["bankroll_after"], 6),
            ])

    # --- Sheet 3: Bankroll Curves ----------------------------------------
    ws3 = wb.create_sheet("Bankroll Curves")
    header = ["window_ts"] + [r["label"] for r in results]
    ws3.append(header)
    n_windows = len(windows)
    for i in range(n_windows):
        row = [windows[i]["window_ts"]]
        for r in results:
            curve = r["curve"]
            row.append(round(curve[i], 6) if i < len(curve) else None)
        ws3.append(row)

    wb.save(output)
    logger.info("Wrote workbook with %d configs to %s", len(results), output)


def _load_candles(path: str) -> list[dict]:
    with open(path) as f:
        return json.load(f)


def _parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Backtest 27 configs and produce an Excel report.")
    src = parser.add_mutually_exclusive_group()
    src.add_argument("--hours", type=float,
                     help="Fetch this many hours of history via backtest.fetch_history.")
    src.add_argument("--candles",
                     help="Path to an offline JSON candle list (no network).")
    parser.add_argument("--output", default="results.xlsx",
                        help="Path for the Excel workbook (default results.xlsx).")
    return parser.parse_args(argv)


def main(argv=None) -> None:
    args = _parse_args(argv)
    starting = _env_float("STARTING_BANKROLL", 1.0)
    min_bet = _env_float("MIN_BET", 1.0)

    if args.candles:
        candles = _load_candles(args.candles)
    else:
        import backtest
        candles = backtest.fetch_history(args.hours if args.hours else 72.0)

    results, windows = run_all(candles, starting, min_bet)
    write_workbook(results, windows, args.output)


if __name__ == "__main__":
    main()
