"""Offline unit tests for compare_runs.py (no network)."""

import json
import math
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import compare_runs  # noqa: E402


def candle(ts, o, h, l, c, v):
    return {"ts": ts, "open": o, "high": h, "low": l, "close": c, "volume": v}


# --- Sizing math per mode --------------------------------------------------

def test_stake_flat():
    assert compare_runs._stake_for("flat", bankroll=10.0, starting=1.0, min_bet=1.0) == 1.0


def test_stake_safe():
    # 25% of bankroll, at least min_bet, never above bankroll.
    assert compare_runs._stake_for("safe", 10.0, 1.0, 1.0) == 2.5
    assert compare_runs._stake_for("safe", 2.0, 1.0, 1.0) == 1.0  # 0.5 -> min_bet


def test_stake_aggressive():
    # Profit above starting; if none, min_bet.
    assert compare_runs._stake_for("aggressive", 10.0, 1.0, 1.0) == 9.0
    assert compare_runs._stake_for("aggressive", 1.0, 1.0, 1.0) == 1.0  # no profit
    assert compare_runs._stake_for("aggressive", 0.5, 1.0, 1.0) == 0.5  # capped at bankroll


# --- Window construction ---------------------------------------------------

def test_build_windows_grouping():
    # 15 one-minute candles -> windows starting at 0, 300k, 600k.
    candles = [candle(i * 60_000, 100 + i, 101 + i, 99 + i, 100 + i, 10)
               for i in range(15)]
    windows = compare_runs.build_windows(candles)
    assert [w["window_ts"] for w in windows] == [0, 300_000, 600_000]
    first = windows[0]
    assert first["window_open"] == candles[0]["open"]
    # resolution close is the window's 5th candle (index 4) close.
    assert first["resolution_close"] == candles[4]["close"]
    # current price is the last closed candle at T-10s (index 3, offset 180s).
    assert first["current_price"] == candles[3]["close"]


def test_build_windows_empty():
    assert compare_runs.build_windows([]) == []


# --- Deterministic PnL / win-loss (mock the strategy signal) ---------------

def _fixed_signal(side, conf=1.0, score=7.0):
    def _analyze(history, window_open, current_price, ticks=None):
        return {"score": score, "confidence": conf, "side": side, "components": {}}
    return _analyze


def test_simulate_config_deterministic_win(monkeypatch):
    monkeypatch.setattr(compare_runs.strategy, "analyze", _fixed_signal("UP"))
    # delta 0 -> entry price 0.50; UP and up_wins True -> win.
    windows = [{
        "window_ts": 0, "window_open": 100.0, "current_price": 100.0,
        "history": [], "resolution_close": 101.0, "up_wins": True,
    }]
    res = compare_runs.simulate_config(windows, threshold=0.0, mode="flat",
                                       starting=1.0, min_bet=1.0)
    assert res["trades"] == 1
    assert res["wins"] == 1
    # stake 1 / entry 0.5 = 2 shares; win pnl = 2*(1-0.5) = 1.0 -> bankroll 2.0
    assert math.isclose(res["final_bankroll"], 2.0)
    assert math.isclose(res["total_pnl"], 1.0)
    assert res["busts"] == 0


def test_simulate_config_deterministic_loss_and_bust(monkeypatch):
    monkeypatch.setattr(compare_runs.strategy, "analyze", _fixed_signal("UP"))
    # UP but up_wins False -> loss of full stake; bankroll goes below min_bet -> bust reset.
    windows = [{
        "window_ts": 0, "window_open": 100.0, "current_price": 100.0,
        "history": [], "resolution_close": 99.0, "up_wins": False,
    }]
    res = compare_runs.simulate_config(windows, threshold=0.0, mode="flat",
                                       starting=1.0, min_bet=1.0)
    assert res["trades"] == 1
    assert res["wins"] == 0
    # lose stake 1 (shares 2 * entry 0.5) -> bankroll 0 < min_bet -> reset to 1.0
    assert res["busts"] == 1
    assert math.isclose(res["final_bankroll"], 1.0)
    assert math.isclose(res["total_pnl"], -1.0)


def test_simulate_config_threshold_skips(monkeypatch):
    monkeypatch.setattr(compare_runs.strategy, "analyze",
                        _fixed_signal("UP", conf=0.2))
    windows = [{
        "window_ts": 0, "window_open": 100.0, "current_price": 100.0,
        "history": [], "resolution_close": 101.0, "up_wins": True,
    }]
    res = compare_runs.simulate_config(windows, threshold=0.5, mode="flat",
                                       starting=1.0, min_bet=1.0)
    assert res["trades"] == 0
    assert math.isclose(res["final_bankroll"], 1.0)
    # curve still records one point per window.
    assert len(res["curve"]) == 1


def test_win_rate_computed(monkeypatch):
    monkeypatch.setattr(compare_runs.strategy, "analyze", _fixed_signal("UP"))
    windows = [
        {"window_ts": 0, "window_open": 100.0, "current_price": 100.0,
         "history": [], "resolution_close": 101.0, "up_wins": True},
        {"window_ts": 300_000, "window_open": 100.0, "current_price": 100.0,
         "history": [], "resolution_close": 99.0, "up_wins": False},
    ]
    res = compare_runs.simulate_config(windows, threshold=0.0, mode="flat",
                                       starting=1.0, min_bet=1.0)
    assert res["trades"] == 2
    assert res["wins"] == 1
    assert math.isclose(res["win_rate"], 0.5)


# --- End-to-end workbook ----------------------------------------------------

def _synthetic_candles(n=200):
    """Deterministic zig-zag walk with enough history and windows."""
    candles = []
    price = 100.0
    for i in range(n):
        # gently oscillating trend so some windows move enough to trade.
        price += math.sin(i / 3.0) * 0.5 + (0.05 if i % 7 else -0.1)
        o = price
        c = price + math.cos(i / 2.0) * 0.3
        h = max(o, c) + 0.2
        low = min(o, c) - 0.2
        v = 10.0 + (5.0 if i % 5 == 0 else 0.0)
        candles.append(candle(i * 60_000, o, h, low, c, v))
    return candles


def test_run_all_produces_27_configs():
    candles = _synthetic_candles()
    results, windows = compare_runs.run_all(candles, starting=1.0, min_bet=1.0)
    assert len(results) == 27
    labels = {r["label"] for r in results}
    assert len(labels) == 27
    assert len(windows) > 0


def test_write_workbook_three_sheets(tmp_path):
    candles = _synthetic_candles()
    results, windows = compare_runs.run_all(candles, starting=1.0, min_bet=1.0)
    out = tmp_path / "results.xlsx"
    compare_runs.write_workbook(results, windows, str(out))
    assert out.exists()
    wb = load_workbook(str(out))
    assert wb.sheetnames == ["Summary", "Best Config Trades", "Bankroll Curves"]
    # Summary: header + 27 config rows.
    ws = wb["Summary"]
    assert ws.max_row == 28
    # Bankroll Curves: header column per config + window_ts.
    curves = wb["Bankroll Curves"]
    assert curves.max_column == 28  # window_ts + 27 configs


def test_main_offline_candles(tmp_path):
    candles = _synthetic_candles()
    candles_path = tmp_path / "candles.json"
    candles_path.write_text(json.dumps(candles))
    out = tmp_path / "out.xlsx"
    compare_runs.main(["--candles", str(candles_path), "--output", str(out)])
    assert out.exists()
    wb = load_workbook(str(out))
    assert wb.sheetnames == ["Summary", "Best Config Trades", "Bankroll Curves"]
